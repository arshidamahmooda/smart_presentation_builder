from pptx import Presentation
from io import BytesIO

def generate_presentation(points):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])

    title = slide.shapes.title
    body = slide.placeholders[1]

    title.text = "Auto-Generated Presentation"
    body.text = ""

    for p in points:
        body.text += f"• {p}\n"

    ppt_stream = BytesIO()
    prs.save(ppt_stream)
    ppt_stream.seek(0)
    return ppt_stream
